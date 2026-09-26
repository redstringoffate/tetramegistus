import json
import os
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.styles.colors import Color

from core.grimoire.styler import apply_grimoire_styles
from core.grimoire.stamper import apply_natal_stamp
from core.astrology.davison import calculate_davison_midpoint

BASE_DIR = os.path.dirname(__file__)
MAPPING_FILE = os.path.abspath(os.path.join(BASE_DIR, '../mappings/a2_ck_mapping.json'))
# 🚀 템플릿 파일명은 a2_ck.xlsx 로 가정합니다.
TEMPLATE_FILE = os.path.abspath(os.path.join(BASE_DIR, '../../../templates_excel/albedo/a2_ck.xlsx'))

def get_safe_float(val, default=0.0):
    try: return float(val) if val is not None else default
    except (ValueError, TypeError): return default

def sanitize_unicode(text):
    if isinstance(text, str):
        return text.replace("\ufe0e", "").replace("\ufe0f", "")
    return text

def compile_a2_ck_grimoire(chart_data, seed_data=None):
    if not os.path.exists(TEMPLATE_FILE):
        raise FileNotFoundError(f"Template missing: {TEMPLATE_FILE}")

    wb = load_workbook(TEMPLATE_FILE)
    base_name = "a2_ck"
    ws = wb[base_name] if base_name in wb.sheetnames else wb.active
    ws.title = base_name

    for s in list(wb.sheetnames):
        if s != base_name: del wb[s]
    wb.active = wb._sheets.index(ws)

    # 1. Albedo (Davison) 데이터 합성 및 Stamper (A2)
    if not seed_data or 'seed1' not in seed_data or 'seed2' not in seed_data:
        raise ValueError("Grimoire requires Albedo seed data. Please run A1 first.")

    s1_raw = seed_data['seed1'].copy()
    s2_raw = seed_data['seed2'].copy()
    dav_hydrated = calculate_davison_midpoint(s1_raw, s2_raw)

    # 🚀 위치 문자열 강제 생성
    d_lat = get_safe_float(dav_hydrated.get("lat"), None)
    d_lng = get_safe_float(dav_hydrated.get("lng"), None)
    
    if d_lat is not None and d_lng is not None:
        coord_str = f"{abs(d_lat):.2f}°{'N' if d_lat>=0 else 'S'}, {abs(d_lng):.2f}°{'E' if d_lng>=0 else 'W'}"
    else:
        coord_str = "Unknown Location"
        
    dav_hydrated["city"] = coord_str
    dav_hydrated["location"] = coord_str
    dav_hydrated["city_name"] = coord_str
    dav_hydrated["location_name"] = coord_str

    apply_natal_stamp(ws, dav_hydrated, method="single", cells=["A2"])

    meta = chart_data.get("metadata", {})
    bodies = chart_data.get("bodies", {})

    with open(MAPPING_FILE, 'r', encoding='utf-8') as f:
        mapping = json.load(f)

    # 2. Metadata (B4, B5, B6)
    meta_map = mapping.get("metadata", {})
    
    day_ref = meta_map.get("day_lord")
    hour_ref = meta_map.get("hour_lord")
    ayan_ref = meta_map.get("ayanamsa")

    if day_ref:
        c_day = ws[day_ref]
        day_val = str(meta.get("day_lord", "-"))
        c_day.value = day_val
        c_day.data_type = 's'
        apply_grimoire_styles(c_day, day_val, skip_color=False)
        color_to_use = copy(c_day.font.color) if c_day.font and c_day.font.color else "000000"
        c_day.font = Font(name="Consolas", size=11, bold=True, color=color_to_use)
        c_day.alignment = Alignment(horizontal='right', vertical='center')

    if hour_ref:
        c_hour = ws[hour_ref]
        hour_val = str(meta.get("hour_lord", "-"))
        c_hour.value = hour_val
        c_hour.data_type = 's'
        apply_grimoire_styles(c_hour, hour_val, skip_color=False)
        color_to_use = copy(c_hour.font.color) if c_hour.font and c_hour.font.color else "000000"
        c_hour.font = Font(name="Consolas", size=11, bold=True, color=color_to_use)
        c_hour.alignment = Alignment(horizontal='right', vertical='center')

    if ayan_ref:
        c_ayan = ws[ayan_ref]
        ayan_val = str(meta.get("ayanamsa", "-")).upper()
        c_ayan.value = ayan_val
        c_ayan.data_type = 's'
        apply_grimoire_styles(c_ayan, ayan_val, skip_color=True)
        color_to_use = copy(c_ayan.font.color) if c_ayan.font and c_ayan.font.color else "000000"
        c_ayan.font = Font(name="Consolas", size=11, bold=False, color=color_to_use)
        c_ayan.alignment = Alignment(horizontal='right', vertical='center')

    # 3. Chara Karaka Data (Rows 10-16)
    layout = mapping.get("layout", {})
    cols = layout.get("cols", {})
    rows = layout.get("rows", {})

    for karaka_key, r_idx in rows.items():
        b_data = bodies.get(karaka_key, {})
        
        info_val = str(b_data.get("info", "-"))
        nak_val = str(b_data.get("nakshatra", "-"))
        g_sa_val = str(b_data.get("graha_sa", "-"))
        g_en_val = str(b_data.get("graha_en", "-"))

        c_info = ws[f"{cols.get('info')}{r_idx}"]
        c_nak = ws[f"{cols.get('nakshatra')}{r_idx}"]
        c_g_sa = ws[f"{cols.get('graha_sa')}{r_idx}"]
        c_g_en = ws[f"{cols.get('graha_en')}{r_idx}"]

        # Information
        c_info.value = info_val
        c_info.data_type = 's'
        apply_grimoire_styles(c_info, sanitize_unicode(info_val), is_info_col=True, skip_color=False)
        c_info.alignment = Alignment(horizontal='left', vertical='center')

        # Nakshatra
        c_nak.value = nak_val
        c_nak.data_type = 's'
        apply_grimoire_styles(c_nak, sanitize_unicode(nak_val), is_info_col=False, skip_color=False)
        c_nak.alignment = Alignment(horizontal='left', vertical='center')

        # Graha (Sanskrit)
        c_g_sa.value = g_sa_val
        c_g_sa.data_type = 's'
        apply_grimoire_styles(c_g_sa, sanitize_unicode(g_sa_val), is_info_col=False, skip_color=False)
        c_g_sa.alignment = Alignment(horizontal='left', vertical='center')

        # Graha (English)
        c_g_en.value = g_en_val
        c_g_en.data_type = 's'
        apply_grimoire_styles(c_g_en, sanitize_unicode(g_en_val), is_info_col=False, skip_color=False)
        c_g_en.alignment = Alignment(horizontal='left', vertical='center')

    # Shrink Logic & 폰트 고정
    for r_idx in range(1, ws.max_row + 1):
        cell_val = ws[f"A{r_idx}"].value
        if str(cell_val).strip() == ".":
            ws.row_dimensions[r_idx].height = 4.5
            for col_char in ["A", "B", "C", "D", "E", "F"]:
                ws[f"{col_char}{r_idx}"].value = ""
        elif r_idx >= 7:
            if ws.row_dimensions[r_idx].height != 4.5:
                ws.row_dimensions[r_idx].height = 16.5

    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                cell.data_type = 's'
            if cell.font:
                safe_color = copy(cell.font.color) if cell.font.color else None
                cell.font = Font(name="Consolas", size=cell.font.size, bold=cell.font.bold, italic=cell.font.italic, color=safe_color)
            else:
                cell.font = Font(name="Consolas")

    # 강제 너비 고정 
    forced_widths = {
        "A": 18.0,  # Karaka
        "B": 35.0,  # Description
        "C": 22.0,  # Information
        "D": 20.0,  # Nakshatra
        "E": 18.0,  # Graha (SA)
        "F": 18.0   # Graha (EN)
    }
    for col_char, target_width in forced_widths.items():
        ws.column_dimensions[col_char].width = target_width

    # A1 셀 병합 해제 방어
    from openpyxl.utils import get_column_letter
    merge_to_remove = None
    for m_range in list(ws.merged_cells.ranges):
        if "A1" in m_range.coord:
            merge_to_remove = m_range
            break
            
    if merge_to_remove:
        max_col_letter = get_column_letter(merge_to_remove.max_col)
        ws.unmerge_cells(merge_to_remove.coord)
        
        ws['B1'].value = ws['A1'].value
        ws['A1'].value = None
        
        if ws['A1'].has_style:
            ws['B1'].font = copy(ws['A1'].font)
            ws['B1'].border = copy(ws['A1'].border)
            ws['B1'].fill = copy(ws['A1'].fill)
            ws['B1'].alignment = copy(ws['A1'].alignment)
            
        ws.merge_cells(f"B1:{max_col_letter}1")

    return wb