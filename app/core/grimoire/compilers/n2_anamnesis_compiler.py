import json
import os
import math
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment

# 🚀 테트라메기스투스 코어 모듈 임포트
from core.grimoire.styler import apply_grimoire_styles
from core.grimoire.stamper import apply_natal_stamp
from core.astrology.engine import calculate_principia, _ensure_float_tz, format_dms_pretty, TROPICAL_SIGNS, SYMBOL_MAP
from core.astrology.divisions.decan import get_decan
from core.astrology.divisions.duad import get_duad
from core.astrology.divisions.dodecatemoria import get_dodecatemoria
from core.astrology.divisions.egyptian_bounds import get_egyptian_bounds
from core.astrology.divisions.sabian_engine import get_sabian_index

BASE_DIR = os.path.dirname(__file__)

MAPPING_FILE = os.path.abspath(os.path.join(BASE_DIR, '../mappings/n2_anamnesis_mapping.json'))
TEMPLATE_FILE = os.path.abspath(os.path.join(BASE_DIR, '../../../templates_excel/nigredo/n2_anamnesis.xlsx'))

SABIAN_PATHS = [
    os.path.abspath(os.path.join(BASE_DIR, '../../../data/render/sabian.json')),
    os.path.abspath(os.path.join(BASE_DIR, '../../../../data/sabian.json')),
    os.path.abspath(os.path.join(BASE_DIR, '../../../static/data/sabian.json')),
    os.path.abspath(os.path.join(BASE_DIR, '../../data/sabian.json')),
]

def get_col(cell_str):
    if not cell_str: return ""
    return "".join([c for c in str(cell_str) if c.isalpha()])

def compile_n2_anamnesis_grimoire(chart_data, seed_data=None):
    if not os.path.exists(TEMPLATE_FILE):
        raise FileNotFoundError(f"Template missing: {TEMPLATE_FILE}")

    wb = load_workbook(TEMPLATE_FILE)
    base_name = "n2_anamnesis"
    ws = wb[base_name] if base_name in wb.sheetnames else wb.active
    ws.title = base_name

    for s in list(wb.sheetnames):
        if s != base_name: 
            del wb[s]
    wb.active = wb._sheets.index(ws)

    meta = chart_data.get("metadata", {})
    system = meta.get("sys_tab", "tropical")
    ayanamsa = meta.get("ayanamsa", "lahiri")
    h_sys = meta.get("h_sys", "P")
    anamnesis_mode = meta.get("anamnesis_mode", "N")

    if not seed_data:
        raise ValueError("Station is vacant. Cannot calculate Anamnesis for Grimoire.")

    date_str = str(seed_data.get('birth_date', '2000-01-01')).split('T')[0]
    time_str = str(seed_data.get('birth_time', '12:00:00'))
    lat = float(seed_data.get("lat", 37.5665))
    lng = float(seed_data.get("lng", 126.9780))
    tz_val = seed_data.get('tz') if seed_data.get('tz') is not None else seed_data.get('timezone', 9.0)
    tz = _ensure_float_tz(tz_val, date_str)
    
    # 아남네시스는 생시 미상일 경우 원천 차단됨
    is_unk = bool(seed_data.get("is_time_unknown", 0))
    if is_unk:
        raise ValueError("Time Unknown. Anamnesis Ritual is locked.")

    apply_natal_stamp(ws, seed_data, method="single", cells=["A2"])

    lang_val = meta.get("language", chart_data.get("language", seed_data.get("language", seed_data.get("lang", "en"))))
    is_ko = "ko" in str(lang_val).lower().strip()

    # 1. 원본 네이탈 연산
    principia_res = calculate_principia(
        date_str=date_str, time_str=time_str, lat=lat, lng=lng, timezone=tz,
        system=system, ayanamsa=ayanamsa, view="zodiac", h_sys=h_sys,
        fixed_star_orb=1.0, is_time_unknown=is_unk
    )

    planets_data = principia_res.get('planets', {})
    lords_data = principia_res.get('lords', {})

    # 2. 🚀 [CORE LOGIC]: Anamnesis Shift & Original House Overlay
    asc_lon = planets_data.get('Ascendant', {}).get('longitude', 0.0)
    ic_lon = planets_data.get('Immum Coeli', {}).get('longitude', 0.0)
    
    # 원본 하우스 커스프 추출 (배경 무대)
    cusps_simple = {int(k): float(v['longitude']) if isinstance(v, dict) else float(v) for k, v in principia_res.get('houses', {}).items()}

    for p_key, p_val in planets_data.items():
        if 'longitude' not in p_val: continue
        old_lon = p_val['longitude']
        
        # 모드별 수학적 시프트
        if anamnesis_mode == 'N': new_lon = (old_lon - asc_lon) % 360
        elif anamnesis_mode == 'R': new_lon = (asc_lon - old_lon) % 360
        elif anamnesis_mode == 'A': new_lon = (old_lon - ic_lon) % 360
        elif anamnesis_mode == 'C': new_lon = (ic_lon - old_lon) % 360
        else: new_lon = old_lon
        
        p_val['longitude'] = new_lon
        p_val['fixed_stars'] = [] # 항성 컨정션 차단
        
        # 하위 속성 재연산
        sign_idx = int(new_lon / 30) % 12
        sign_name = TROPICAL_SIGNS[sign_idx]
        deg_in_sign = new_lon % 30
        
        p_val['dms'] = format_dms_pretty(new_lon)
        p_val['sign'] = sign_idx
        p_val['is_anaretic'] = (deg_in_sign >= 29.0)
        p_val['duad'] = SYMBOL_MAP.get(get_duad(sign_name, deg_in_sign), "-")
        p_val['dodeca'] = SYMBOL_MAP.get(get_dodecatemoria(deg_in_sign), "-")
        p_val['decan'] = SYMBOL_MAP.get(get_decan(sign_name, deg_in_sign), "-")
        p_val['bound'] = SYMBOL_MAP.get(get_egyptian_bounds(sign_name, deg_in_sign), "-")
        p_val['sabian_index'] = get_sabian_index(new_lon)

        # 재귀된 영혼이 네이탈 차트의 어느 하우스 무대에 떨어지는지 판별
        if len(cusps_simple) == 12:
            found_house = 1
            for h_num in range(1, 13):
                cur = cusps_simple[h_num]
                nxt = cusps_simple[h_num+1] if h_num < 12 else cusps_simple[1]
                if cur < nxt:
                    if cur <= new_lon < nxt: 
                        found_house = h_num
                        break
                else:
                    if cur <= new_lon < 360 or 0 <= new_lon < nxt: 
                        found_house = h_num
                        break
            p_val['house'] = str(found_house)

    # 3. 사비안 심볼 로더
    sabian_dict = {}
    for spath in SABIAN_PATHS:
        if os.path.exists(spath):
            try:
                with open(spath, 'r', encoding='utf-8') as f: sabian_dict = json.load(f)
                break
            except: pass
    
    def get_sabian_text(idx):
        if not idx: return ""
        idx_str = str(idx)
        def extract_text(entry):
            if isinstance(entry, str): return entry
            if isinstance(entry, dict):
                ko_keys = ["text_ko", "ko", "desc_ko"]
                en_keys = ["text_en", "en", "desc_en"]
                primary_keys = ko_keys if is_ko else en_keys
                for k in primary_keys + (en_keys if is_ko else ko_keys) + ["text", "desc"]:
                    if k in entry and entry[k]: return str(entry[k])
            return ""
        if isinstance(sabian_dict, dict):
            if idx_str in sabian_dict: return extract_text(sabian_dict[idx_str])
        elif isinstance(sabian_dict, list):
            for item in sabian_dict:
                if isinstance(item, dict) and str(item.get("index", "")) == idx_str: return extract_text(item)
        return ""

    with open(MAPPING_FILE, 'r', encoding='utf-8') as f:
        mapping = json.load(f)

    # 4. 메타데이터 인젝션
    day_val = lords_data.get("day", "")
    hour_val = lords_data.get("hour", "")
    mode_names = {"N": "Nigredo", "A": "Albedo", "C": "Citrinitas", "R": "Rubedo"}
    
    ws[mapping["metadata"]["day_lord"]] = day_val
    ws[mapping["metadata"]["hour_lord"]] = hour_val
    ws[mapping["metadata"]["anamnesis_mode"]] = mode_names.get(anamnesis_mode, "Nigredo")
    
    house_sys_map = {"P": "PLACIDUS", "W": "WHOLE SIGN", "K": "KOCH"}
    h_sys_full = house_sys_map.get(str(h_sys).upper().strip(), str(h_sys).upper())
    ws[mapping["metadata"]["house_sys"]] = h_sys_full

    for ref in ["day_lord", "hour_lord", "anamnesis_mode", "house_sys"]:
        cell_ref = mapping["metadata"][ref]
        apply_grimoire_styles(ws[cell_ref], ws[cell_ref].value, skip_color=("|" in str(ws[cell_ref].value)))

    # 5. 행성 및 포인트 데이터 루프 (별 삭제됨)
    all_bodies = []
    for category in ["planets", "asteroids", "lilith_nodes", "fates"]: 
        for body_name, map_data in mapping.get(category, {}).items():
            engine_name = body_name
            if engine_name == "North Node (t)": engine_name = "North Node (t)"
            elif engine_name == "North Node (m)": engine_name = "Rahu"
            elif engine_name == "South Node (m)": engine_name = "Ketu"
            elif engine_name == "South Node (t)": engine_name = "South Node (t)"
            
            p_data = planets_data.get(engine_name)
            if p_data and isinstance(p_data, dict): 
                all_bodies.append({ "name": body_name, "map": map_data, "data": p_data })
    
    day_lords_list = [d.strip() for d in str(day_val).split('|')]

    for body in all_bodies:
        b_map = body["map"]
        p_data = body["data"] 
        body_name = body["name"]
        row_idx = b_map.get("row_start", 0)
        if not row_idx: continue

        col_info = get_col(b_map.get("info", "B"))
        col_ruler = get_col(b_map.get("ruler", "C"))
        col_dignity = get_col(b_map.get("dignity", "D"))
        col_house = get_col(b_map.get("house", "E"))
        col_duad = get_col(b_map.get("duad", "F"))
        col_dodeca = get_col(b_map.get("dodeca", "G"))
        col_decan = get_col(b_map.get("decan", "H"))
        col_bounds = get_col(b_map.get("bounds", "I"))
        col_sabian = get_col(b_map.get("sabian", "J"))
        
        info_text = p_data.get("dms", "")
        if body_name == hour_val: info_text = info_text.upper()
            
        ws[f"{col_info}{row_idx}"] = info_text
        ws[f"{col_ruler}{row_idx}"] = str(p_data.get("ruler", ""))
        ws[f"{col_dignity}{row_idx}"] = str(p_data.get("dignity", "None")) or "None"
        ws[f"{col_house}{row_idx}"] = str(p_data.get("house", ""))
        ws[f"{col_duad}{row_idx}"] = str(p_data.get("duad", ""))
        ws[f"{col_dodeca}{row_idx}"] = str(p_data.get("dodeca", ""))
        ws[f"{col_decan}{row_idx}"] = str(p_data.get("decan", ""))
        ws[f"{col_bounds}{row_idx}"] = str(p_data.get("bound", ""))
        
        # 🚀 [완벽 수복]: sabian.json은 0번부터 시작하므로 +1 절대 금지.
        # 또한 인덱스가 '0'일 때 파이썬이 빈 값(Falsy)으로 착각하지 않도록 명시적 처리.
        s_idx = None
        if "longitude" in p_data:
            try:
                lon_val = float(p_data["longitude"]) % 360
                s_idx = int(math.floor(lon_val)) 
            except:
                s_idx = p_data.get("sabian_index")
        else:
            s_idx = p_data.get("sabian_index")
            
        s_val_str = str(s_idx).strip() if s_idx is not None else ""
        sabian_text = ""
        
        if s_val_str.isdigit(): sabian_text = get_sabian_text(s_val_str)
        if not sabian_text:
            fallback_txt = p_data.get("sabian", p_data.get("sabian_text", p_data.get("sabian_symbol", "")))
            if isinstance(fallback_txt, dict):
                if is_ko: sabian_text = fallback_txt.get("text_ko", fallback_txt.get("ko", fallback_txt.get("text_en", fallback_txt.get("en", ""))))
                else: sabian_text = fallback_txt.get("text_en", fallback_txt.get("en", fallback_txt.get("text_ko", fallback_txt.get("ko", ""))))
            elif isinstance(fallback_txt, str): sabian_text = fallback_txt

        if not sabian_text:
            frontend_sabian = chart_data.get("bodies", {}).get(body_name, {}).get("sabian", "")
            if frontend_sabian: sabian_text = frontend_sabian

        if not sabian_text and s_val_str.isdigit():
            sabian_text = f"[{s_val_str}°] Symbol rendering fallback"

        ws[f"{col_sabian}{row_idx}"] = sabian_text

        for c_char in "BCDEFGHIJ":
            target_cell = ws[f"{c_char}{row_idx}"]
            apply_grimoire_styles(
                target_cell, 
                target_cell.value, 
                is_info_col=(c_char == col_info), 
                skip_color=(c_char == col_sabian),
                is_day_lord=(body_name in day_lords_list),   
                is_hour_lord=(body_name == hour_val), 
                is_anaretic=p_data.get("is_anaretic", False)
            )
        ws[f"{col_house}{row_idx}"].alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows():
        for cell in row:
            if cell.font:
                cell.font = Font(name="Consolas", size=cell.font.size, bold=cell.font.bold, italic=cell.font.italic, color=cell.font.color)

    return wb